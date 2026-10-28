# true_labels.py
# Loads ground-truth Bloom's Taxonomy labels from the selected worksheet in
# the local Excel workbook.
#
# Exported symbols
# ----------------
# true_labels_3   : list[int]  – 3-class labels matching n_clusters=3
#                                0=LOTS, 1=MOTS, 2=HOTS
# true_labels_6   : list[int]  – 6-class Bloom's level labels
#                                0=Knowledge, 1=Comprehension, 2=Application,
#                                3=Analysis,  4=Synthesis,    5=Evaluation
# label_names_3   : list[str]  – human-readable names for the 3-class labels
# label_names_6   : list[str]  – human-readable names for the 6-class labels

import os

from openpyxl import load_workbook
from preprocessing import _get_selected_dataset_name

dataset_name = _get_selected_dataset_name()

# Read the same non-empty question rows as preprocessing.py, with LEVEL in F
# and BT LEVEL in G, so labels remain aligned with the embeddings.
workbook_path = os.path.join(os.path.dirname(__file__), '20_01_2022_dataset.xlsx')
workbook = load_workbook(workbook_path, read_only=True, data_only=True)
try:
    if dataset_name not in workbook.sheetnames:
        raise ValueError(f"Worksheet '{dataset_name}' not found in '{workbook_path}'.")
    worksheet = workbook[dataset_name]
    label_rows = [
        (level, bt_level)
        for question, level, bt_level in worksheet.iter_rows(
            min_row=2, min_col=5, max_col=7, values_only=True
        )
        if question is not None and str(question).strip()
    ]
finally:
    workbook.close()

raw_level = [level if level is not None else '' for level, _ in label_rows]
raw_bt_level = [bt_level if bt_level is not None else '' for _, bt_level in label_rows]

# ── 3-class mapping: LOTS=0, MOTS=1, HOTS=2 ──────────────────────────────────
_level_map = {'lots': 0, 'mots': 1, 'hots': 2}
label_names_3 = ['LOTS', 'MOTS', 'HOTS']

true_labels_3 = [_level_map.get(str(v).strip().lower(), -1) for v in raw_level]

# ── 6-class mapping: normalise known variants ─────────────────────────────────
_bt_map = {
    'knowledge':     0,
    'comprehension': 1,
    'application':   2,
    'apply':         2,   # variant found in FBF sheet
    'analysis':      3,
    'anaysis':       3,   # typo found in FAS sheet
    'synthesis':     4,
    'evaluation':    5,
    'evaluate':      5,   # variant found in FBF sheet
}
label_names_6 = ['Knowledge', 'Comprehension', 'Application',
                  'Analysis', 'Synthesis', 'Evaluation']

true_labels_6 = [_bt_map.get(str(v).strip().lower(), -1) for v in raw_bt_level]

# ── Sanity check ──────────────────────────────────────────────────────────────
_unmapped_3 = [v for v, l in zip(raw_level,    true_labels_3) if l == -1]
_unmapped_6 = [v for v, l in zip(raw_bt_level, true_labels_6) if l == -1]
if _unmapped_3:
    print(f"[true_labels] WARNING: unmapped LEVEL values: {set(_unmapped_3)}")
if _unmapped_6:
    print(f"[true_labels] WARNING: unmapped BT LEVEL values: {set(_unmapped_6)}")